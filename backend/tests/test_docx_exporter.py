"""Word（.docx）导出：字节可读、版式复用 ResumeLayout、分页与 PDF ±1。"""
from io import BytesIO

import pytest
from docx import Document

from app.schemas.resume import ResumeContent
from app.services.docx_exporter import build_resume_docx
from app.services.pdf_exporter import build_resume_pdf, font_available
from app.services.resume.resume_sample import sample_resume_content

needs_font = pytest.mark.skipif(not font_available(), reason="本机没有可用的中文字体")


def _resume() -> ResumeContent:
    return ResumeContent(
        name="张三",
        phone="13800000000",
        summary="负责后端服务的设计与开发。",
        experience=[{"company": "字节跳动", "role": "工程师", "description": ["实现检索接口"]}],
        skills=[{"name": "Python", "level": "熟练"}],
    )


def test_docx_builds_a_valid_word_document():
    result = build_resume_docx(_resume(), template="classic", page_limit=1)

    # .docx 是 ZIP 容器，文件头以 PK 开头；且能被 python-docx 重新打开。
    assert result.content.startswith(b"PK")
    document = Document(BytesIO(result.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "张三" in text
    assert "字节跳动" in text
    assert "实现检索接口" in text


def test_docx_reuses_layout_margins_and_respects_margin_override():
    result = build_resume_docx(_resume(), template="classic", page_limit=1, margin_mm=20)

    document = Document(BytesIO(result.content))
    section = document.sections[0]
    # 页边距直接覆盖 20mm，且四边一致（复用 ResumeLayout.margin_mm 的同一来源）。
    assert round(section.top_margin.mm, 2) == pytest.approx(20.0, abs=0.01)
    assert round(section.left_margin.mm, 2) == pytest.approx(20.0, abs=0.01)
    assert round(section.right_margin.mm, 2) == pytest.approx(20.0, abs=0.01)
    assert round(section.bottom_margin.mm, 2) == pytest.approx(20.0, abs=0.01)


def test_docx_uses_template_default_margin_when_not_overridden():
    from app.services.resume.resume_templates import TEMPLATE_LAYOUT_DEFAULTS

    expected = float(TEMPLATE_LAYOUT_DEFAULTS["classic"]["padding_mm"])
    result = build_resume_docx(_resume(), template="classic", page_limit=1)

    document = Document(BytesIO(result.content))
    section = document.sections[0]
    assert round(section.top_margin.mm, 2) == pytest.approx(expected, abs=0.01)


def test_docx_excludes_photo_when_include_photo_is_false():
    # 用一张合法 PNG 的 data URL；关闭照片后文档里不应出现任何图片。
    png = (
        "data:image/png;base64,"
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
        "YAAAAAYAAjCB0C8AAAAASUVORK5CYII="
    )
    resume = ResumeContent(name="张三", photo=png, summary="有照片。")

    with_photo = build_resume_docx(resume, template="classic", page_limit=1, include_photo=True)
    without_photo = build_resume_docx(resume, template="classic", page_limit=1, include_photo=False)

    # 带照片的文档里至少有 1 张图片（inline shape），关掉后没有。
    doc_with = Document(BytesIO(with_photo.content))
    doc_without = Document(BytesIO(without_photo.content))
    assert len(doc_with.inline_shapes) >= 1
    assert len(doc_without.inline_shapes) == 0


def test_docx_gender_is_a_separate_muted_run_not_the_name_run():
    """Word 里性别必须是独立的 muted 灰 run，不再与姓名拼成一个强调色 run。

    此前 `name_line = name + "　" + gender` 拼成一个 run，性别跟着姓名一起变大变粗、
    变强调色。拆分后姓名 run 是粗体 + 强调色 + name_ratio 字号，性别 run 是常规 + muted +
    name_extra_ratio 字号。
    """
    from app.services.resume.resume_templates import TEMPLATE_LAYOUT_DEFAULTS

    resume = ResumeContent(name="张三", gender="男", summary="一句话总结。")
    result = build_resume_docx(resume, template="classic", page_limit=1)
    doc = Document(BytesIO(result.content))

    # 第一个段落应是页头照片段或姓名段；找到含"张三"的段落。
    name_paragraph = next(p for p in doc.paragraphs if "张三" in p.text)
    runs = name_paragraph.runs
    assert len(runs) >= 2, "姓名与性别应分成两个 run"
    name_run = next(r for r in runs if "张三" in r.text)
    gender_run = next(r for r in runs if "男" in r.text)

    defaults = TEMPLATE_LAYOUT_DEFAULTS["classic"]
    base_px = 14.0
    # 姓名：粗体、强调色、name_ratio 字号
    assert name_run.font.bold is True
    # python-docx 把 Pt 量化到 0.5pt（19.53 → 19.5），用 0.05 容差。
    assert name_run.font.size.pt == pytest.approx(
        base_px * float(defaults["name_ratio"]) * 0.75, abs=0.05
    )
    # 性别：非常规粗体（None 或 False）、name_extra_ratio 字号、且字号小于姓名
    assert not gender_run.font.bold
    assert gender_run.font.size.pt == pytest.approx(
        base_px * float(defaults["name_extra_ratio"]) * 0.75, abs=0.05
    )
    assert gender_run.font.size.pt < name_run.font.size.pt


def test_docx_photo_is_cover_cropped_before_embedding():
    """Word 的照片也应先按模板照片框比例做 object-fit: cover 裁剪，不变形。

    用一张横图（宽高比 ≠ classic 照片框 0.7875）：裁剪后嵌入的图片宽度仍是模板的
    `photo_width_ratio`（由 layout 给定），高度随裁后比例走——这条断言"照片真的进了文档
    且没让导出崩"。几何精确比对在 pdf_exporter 的纯函数用例里已钉。
    """
    import base64
    from io import BytesIO

    from PIL import Image

    from app.services.resume.resume_templates import TEMPLATE_LAYOUT_DEFAULTS

    # 用 Pillow 现造一张 4x1 的横图 PNG（与 classic 照片框 0.7875 比例明显不同）。
    buf = BytesIO()
    Image.new("RGB", (400, 100), (0, 128, 200)).save(buf, format="PNG")
    png = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode("ascii")
    resume = ResumeContent(name="张三", photo=png, summary="带照片。")
    result = build_resume_docx(resume, template="classic", page_limit=1, include_photo=True)
    doc = Document(BytesIO(result.content))
    assert len(doc.inline_shapes) >= 1
    # 嵌入宽度应等于模板照片框宽度（mm 换算成 EMU）
    defaults = TEMPLATE_LAYOUT_DEFAULTS["classic"]
    expected_width_emu = int(round(14.0 * float(defaults["photo_width_ratio"]) * 0.264583 * 36000))
    assert abs(doc.inline_shapes[0].width - expected_width_emu) < expected_width_emu * 0.1


@needs_font
def test_docx_page_estimate_is_within_one_of_pdf():
    resume = sample_resume_content()
    pdf = build_resume_pdf(resume, template="classic", page_limit=2, font_scale="standard")
    docx = build_resume_docx(resume, template="classic", page_limit=2, font_scale="standard")

    assert docx.pages is not None
    assert abs(docx.pages - pdf.pages) <= 1
